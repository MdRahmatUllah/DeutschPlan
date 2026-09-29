"""Writes a meaning language's staged content into a workbook's columns (#1080, #1082-#1084).

    python tools/merge_language_columns.py data/German_A1_Tracker.xlsx en ru pl \\
        --out data/_staging/pilot/German_A1_Tracker.xlsx

Each code's staging is `<workbook folder>/_staging/<code>/<workbook stem>.json`
(`--staging` names another root):

    {"language": "Russian",
     "words":   [{"key": [German, POS, English],
                  "Meaning": "...", "Pronunciation": "...", "Examples": ["line 1", "line 2"]}],
     "grammar": [{"key": [Level, Topic],
                  "Topic": "...", "Rule": "...", "Example": "...", "Watch out": "..."}]}

A word is found by its German, POS and English cells, a grammar topic by its
Level and Topic cells; each must be one row of the workbook. Every entry of a
part carries the same fields, and a field becomes the column `<field> (<language>)`
of #1080's convention: All Words `Meaning`, `Pronunciation`, `Examples` (one line
per line of `Examples (DE)`, newline-joined), Grammar `Topic`, `Rule`, `Example`,
`Watch out`. A new column goes after the sheet's last one, its header styled like
the English column's; a column that is there already (a second merge after a
review) is rewritten in place. The staging is the whole column: a row it lacks
is left blank.

The output is a new file: the input is never written. openpyxl rewrites the
Dashboard chart on save (and older versions drop it), so the original chart and
drawing parts are copied back. Then the output is checked against the input:
every cell outside the language columns is unchanged (value and style), and
the chart parts are the original bytes. A failed check deletes the output.
"""
from __future__ import annotations

import argparse
import json
import sys
import zipfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles.cell_style import StyleArray
from openpyxl.chart import LineChart
from openpyxl.utils import get_column_letter

WORDS, GRAMMAR = "All Words", "Grammar"
#: where a sheet's staging is, the cells that key a row, and each field's
#: English column (whose header and cell styles it copies) and width
SHEETS = {
    WORDS: ("words", ("German", "POS", "English"),
            {"Meaning": ("English", 26), "Pronunciation": ("Pronunciation (Bangla)", 22),
             "Examples": ("Examples (EN)", 38)}),
    GRAMMAR: ("grammar", ("Level", "Topic"),
              {"Topic": ("Topic", 32), "Rule": ("Rule", 64), "Example": ("Example (EN)", 40),
               "Watch out": ("Watch out", 42)}),
}
CHART_PARTS = ("xl/charts/", "xl/drawings/")


class MergeError(Exception):
    pass


def header_row(ws, key: tuple[str, ...]) -> tuple[int, dict[str, int]]:
    """The header row and its columns by header: the first row holding every key header."""
    for r in range(1, min(ws.max_row, 10) + 1):
        cols = {c.value: c.column for c in ws[r] if c.value is not None}
        if all(k in cols for k in key):
            return r, cols
    raise MergeError(f"{ws.title}: no header row with {', '.join(key)}")


def rows_by_key(ws, hrow: int, cols: dict[str, int], key: tuple[str, ...]) -> dict[tuple, int]:
    rows: dict[tuple, int] = {}
    for r in range(hrow + 1, ws.max_row + 1):
        k = tuple(ws.cell(r, cols[name]).value for name in key)
        if all(v is None for v in k):
            continue
        if k in rows:
            raise MergeError(f"{ws.title}: rows {rows[k]} and {r} have the same key {list(k)}")
        rows[k] = r
    return rows


def ungroup(ws) -> None:
    """openpyxl keeps a `<col min max>` group as one entry: split it, so a new
    column's width leaves its neighbours' alone."""
    for dim in list(ws.column_dimensions.values()):
        if dim.min and dim.max and dim.max > dim.min:
            for i in range(dim.min + 1, dim.max + 1):
                one = copy(dim)
                one.index, one.min, one.max = get_column_letter(i), i, i
                ws.column_dimensions[one.index] = one
            dim.max = dim.min


def merge_sheet(ws, language: str, entries: list[dict], key: tuple[str, ...],
                fields: dict[str, tuple[str, int]]) -> tuple[set[int], int]:
    """Writes one language's part; returns the columns written and the rows filled."""
    hrow, cols = header_row(ws, key)
    rows = rows_by_key(ws, hrow, cols, key)
    present = set(entries[0]) - {"key"}
    by_row: dict[int, dict] = {}
    for e in entries:
        if set(e) - {"key"} != present:
            raise MergeError(f"{language} {ws.title} {e['key']}: fields {sorted(set(e) - {'key'})}, "
                             f"the first entry has {sorted(present)}")
        k = tuple(e["key"])
        if k not in rows:
            raise MergeError(f"{language} {ws.title}: no row has the key {list(k)}")
        if rows[k] in by_row:
            raise MergeError(f"{language} {ws.title}: {list(k)} is staged twice")
        by_row[rows[k]] = e
    if present - set(fields):
        raise MergeError(f"{language} {ws.title}: unknown fields {sorted(present - set(fields))}")
    ungroup(ws)
    written: set[int] = set()
    for field, (like, width) in fields.items():
        if field not in present:
            continue
        header = f"{field} ({language})"
        col = cols.get(header)
        if col is None:
            col = cols[header] = max(cols.values()) + 1
            # the check skips the written columns, so a cell there would be lost unseen
            taken = [r for r in range(1, ws.max_row + 1) if ws.cell(r, col).value is not None]
            if taken:
                raise MergeError(f"{language} {ws.title}: the new column {header} would overwrite "
                                 f"{get_column_letter(col)}{taken[0]}, which has no header but a value")
            ws.cell(hrow, col).value = header
            ws.cell(hrow, col)._style = copy(ws.cell(hrow, cols[like])._style)
            ws.column_dimensions[get_column_letter(col)].width = width
        written.add(col)
        for r in rows.values():
            e = by_row.get(r)
            value = None if e is None else e[field]
            if field == "Examples" and e is not None:
                german = str(ws.cell(r, cols["Examples (DE)"]).value or "").split("\n")
                if len(value) != len(german):
                    raise MergeError(f"{language} {e['key']}: {len(value)} example lines, "
                                     f"the German has {len(german)}")
                value = "\n".join(value)
            ws.cell(r, col).value = value
            ws.cell(r, col)._style = copy(ws.cell(r, cols[like])._style)
    return written, len(by_row)


def chart_parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist() if n.startswith(CHART_PARTS)}


def put_parts(path: Path, parts: dict[str, bytes]) -> None:
    """Replaces these parts of the xlsx at [path] (as build_trackers.py's swap_chart)."""
    with zipfile.ZipFile(path) as z:
        allparts = {i.filename: z.read(i.filename) for i in z.infolist()}
    if {n for n in allparts if n.startswith(CHART_PARTS)} != set(parts):
        raise MergeError(f"chart parts: the output has {sorted(n for n in allparts if n.startswith(CHART_PARTS))}, "
                         f"the input {sorted(parts)}")
    allparts.update(parts)
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in allparts.items():
            z.writestr(name, data)
    tmp.replace(path)


def check(src: Path, out: Path, written: dict[str, set[int]]) -> None:
    """Every cell outside the written columns is as in [src], and so are the chart parts."""
    blank = (None, tuple(StyleArray()))

    def fp(cell):
        # an ArrayFormula compares by identity: compare its fields
        return blank if cell is None else (getattr(cell.value, "__dict__", cell.value), tuple(cell._style))

    a, b = load_workbook(src), load_workbook(out)
    if a.sheetnames != b.sheetnames:
        raise MergeError(f"check: sheets {b.sheetnames}, the input has {a.sheetnames}")
    changed = []
    for ws in a.worksheets:
        other, skip = b[ws.title], written.get(ws.title, set())
        for rc in set(ws._cells) | set(other._cells):
            if rc[1] not in skip and fp(ws._cells.get(rc)) != fp(other._cells.get(rc)):
                changed.append(f"{ws.title}!{get_column_letter(rc[1])}{rc[0]}")
    if changed:
        raise MergeError(f"check: {len(changed)} cells outside the new columns changed: {sorted(changed)[:10]}")
    if chart_parts(src) != chart_parts(out):
        raise MergeError("check: the chart parts differ from the input's")


def merge(src: Path, out: Path, stagings: list[dict]) -> list[str]:
    """Merges [stagings] into a copy of [src] at [out]; returns a report line per language and part."""
    if out.resolve() == src.resolve():
        raise MergeError("--out is the input: the merge never writes its input")
    wb = load_workbook(src)
    written: dict[str, set[int]] = {}
    report = []
    for st in stagings:
        for sheet, (part, key, fields) in SHEETS.items():
            if st.get(part):
                cols, n = merge_sheet(wb[sheet], st["language"], st[part], key, fields)
                written.setdefault(sheet, set()).update(cols)
                report.append(f"{st['language']}: {sheet} {n} rows")
    dash = wb["Dashboard"] if "Dashboard" in wb.sheetnames else None
    if dash is not None and not dash._charts:
        # ponytail: openpyxl 3.1.5 reads the chart back; one that drops it gets a
        # placeholder here, and the original chart and drawing XML replace it below
        dash.add_chart(LineChart(), "L2")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    try:
        put_parts(out, chart_parts(src))
        check(src, out, written)
    except Exception:
        out.unlink(missing_ok=True)
        raise
    return report


def main(argv: list[str] | None = None) -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument("workbook", type=Path)
    p.add_argument("codes", nargs="+", help="language codes, e.g. en ru pl")
    p.add_argument("--out", type=Path, required=True, help="the merged copy (never the input)")
    p.add_argument("--staging", type=Path, help="staging root (default: <workbook folder>/_staging)")
    a = p.parse_args(argv)
    root = a.staging or a.workbook.parent / "_staging"
    stagings = [json.loads((root / c / f"{a.workbook.stem}.json").read_text(encoding="utf-8"))
                for c in a.codes]
    try:
        report = merge(a.workbook, a.out, stagings)
    except MergeError as e:
        print(f"error: {e}", file=sys.stderr)
        return 1
    print("\n".join(report + [f"wrote {a.out}; check passed: other cells and the chart unchanged"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
