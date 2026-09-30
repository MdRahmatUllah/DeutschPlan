"""#1080: meaning languages, read from the workbooks' columns, shipped at 100 %.

A language is its columns: `Meaning (Russian)`, `Pronunciation (Russian)`,
`Examples (Russian)`, and the Grammar's `Topic (Russian)` … `Watch out
(Russian)`. The fixture workbooks carry a whole Russian when asked to
(`write_all(dir, ("Russian",))`), which is how a new language is proved to
need no code.
"""

from __future__ import annotations

import sqlite3
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from excel_to_sqlite import (  # noqa: E402
    PipelineError,
    main as build_main,
    read_manifest,
    read_sources,
    read_workbook,
)
from fixtures.make_workbooks import BOOK_LEVELS, write_all, write_workbook  # noqa: E402


@pytest.fixture(scope="module")
def russian(tmp_path_factory) -> Path:
    """The six fixture workbooks, each with a complete Russian."""
    directory = tmp_path_factory.mktemp("russian")
    write_all(directory, ("Russian",))
    return directory


def manifest(directory: Path, books: dict[str, Path] | None = None, tips: Path | None = None) -> Path:
    entries = {"workbooks": [{"file": str((books or {}).get(name, directory / name))} for name in BOOK_LEVELS]}
    if tips:
        entries["tips"] = str(tips)
    path = directory / "manifest.yaml"
    path.write_text(yaml.safe_dump(entries), encoding="utf-8")
    return path


def build(tmp_path: Path, manifest_path: Path, *flags: str) -> sqlite3.Connection:
    out = tmp_path / "build" / "content.db"
    argv = ["--manifest", str(manifest_path), "--out", str(out), "--previous", str(tmp_path / "none"), *flags]
    assert build_main(argv) == 0
    return sqlite3.connect(out)


def copy_with(path: Path, tmp_path: Path, mutate) -> Path:
    book = load_workbook(path)
    mutate(book)
    out = tmp_path / f"{path.stem}-{mutate.__name__}.xlsx"
    book.save(out)
    return out


def column_of(sheet, header: str) -> int:
    for row in sheet.iter_rows(min_row=1, max_row=5):
        for cell in row:
            if cell.value == header:
                return cell.column
    raise AssertionError(f"no {header!r}")


def langs(db: sqlite3.Connection, table: str) -> dict[str, int]:
    return dict(db.execute(f"SELECT lang, count(*) FROM {table} GROUP BY lang"))


class TestReading:
    def test_1080_a_languages_columns_are_read_by_pattern(self, russian):
        source = read_workbook(russian / "German_A1_Tracker.xlsx")
        word = source.words[0]
        assert word.texts[("meaning", "ru")] == f"[Russian meaning] {word.english}"
        assert word.texts[("pronunciation", "ru")].startswith("[Russian pronunciation]")
        assert word.texts[("examples", "ru")].splitlines() == [
            f"[Russian] {line}" for line in word.examples_de.splitlines()
        ]
        topic = source.grammar[0]
        assert set(topic.texts) >= {("topic", "ru"), ("rule", "ru"), ("example", "ru")}
        assert source.unmatched == {"All Words": [], "Grammar": []}

    def test_1080_todays_headers_are_english_and_bangla_under_either_name(self, russian, tmp_path):
        renames = {
            "English": "Meaning (English)",
            "Bangla meaning": "Meaning (Bangla)",
            "Examples (EN)": "Examples (English)",
            "Topic": "Topic (English)",
            "Rule": "Rule (English)",
            "Example (EN)": "Example (English)",
            "Watch out": "Watch out (English)",
        }

        def new_names(book):
            for sheet in ("All Words", "Grammar"):
                for row in book[sheet].iter_rows(min_row=1, max_row=3):
                    for cell in row:
                        cell.value = renames.get(cell.value, cell.value)

        path = russian / "German_B1_Tracker.xlsx"
        before, after = read_workbook(path), read_workbook(copy_with(path, tmp_path, new_names))
        assert after.unmatched == {"All Words": [], "Grammar": []}
        fields = ("english", "bangla", "pron_bn", "examples_en")
        assert [[getattr(w, f) for f in fields] for w in after.words] == [
            [getattr(w, f) for f in fields] for w in before.words
        ]
        grammar = ("topic", "rule", "example_en", "watch_out")
        assert [[getattr(r, f) for f in grammar] for r in after.grammar] == [
            [getattr(r, f) for f in grammar] for r in before.grammar
        ]

    def test_1080_a_language_name_the_pipeline_does_not_know_stops_the_build(self, tmp_path):
        path = tmp_path / "German_A1_Tracker.xlsx"
        write_workbook(path, ["A1"], ("Klingon",))
        with pytest.raises(PipelineError, match=r"Meaning \(Klingon\) names no language.*English, Bangla, Russian"):
            read_workbook(path)

    def test_1080_a_language_missing_a_column_in_one_workbook_stops_the_build(self, russian, tmp_path):
        # C2 without its Russian: every other workbook has it.
        plain = tmp_path / "German_C2_Tracker.xlsx"
        write_workbook(plain, ["C2"])
        with pytest.raises(PipelineError) as refused:
            read_sources(read_manifest(manifest(russian, {"German_C2_Tracker.xlsx": plain})))
        message = str(refused.value)
        assert "German_C2_Tracker.xlsx!All Words has no 'Examples (Russian)', 'Meaning (Russian)', 'Pronunciation (Russian)' columns" in message
        assert "German_C2_Tracker.xlsx!Grammar has no" in message and "'Topic (Russian)'" in message


class TestTheGate:
    def test_1080_a_complete_language_ships_into_every_table(self, russian, tmp_path, capsys):
        tips = tmp_path / "tips.csv"
        tips.write_text(
            "match_type,match,tip_en,tip_bn,tip_ru,tags\n"
            "pattern,^Tisch,A tip.,একটি টিপ।,Совет.,\n"
            "pattern,^Lampe,,,Только по-русски.,\n",
            encoding="utf-8",
        )
        db = build(tmp_path, manifest(russian, tips=tips))
        err = capsys.readouterr().err
        assert "language Russian (ru): meanings 180/180, pronunciations 180/180, examples 360/360, grammar topics 36/36" in err
        assert [row[0] for row in db.execute("SELECT code FROM course_languages ORDER BY ord")] == ["en", "bn", "ru"]
        assert db.execute("SELECT own_name, script FROM course_languages WHERE code = 'ru'").fetchone() == ("Русский", "Cyrl")
        # English's and Bangla's texts are the course's own columns, not rows
        # (#1096); English has no pronunciation guide here to need one.
        assert langs(db, "word_meanings") == {"ru": 180}
        # Bangla has no examples or grammar columns: its learners read English's.
        assert langs(db, "word_example_translations") == {"ru": 360}
        assert langs(db, "grammar_translations") == {"ru": 36}
        tip_rows = langs(db, "word_tips")
        english, bangla = db.execute(
            "SELECT count(tip_en), count(tip_bn) FROM interference_tips"
        ).fetchone()
        assert set(tip_rows) == {"ru"} and tip_rows["ru"] > english == bangla > 0

    def test_1096_english_and_bangla_are_the_courses_own_columns_not_rows(self, russian, tmp_path):
        db = build(tmp_path, manifest(russian))
        for table in ("word_meanings", "word_example_translations", "grammar_translations", "word_tips"):
            assert {"en", "bn"}.isdisjoint(langs(db, table)), table
        # Both still ship, and their texts are the columns'.
        assert [row[0] for row in db.execute("SELECT code FROM course_languages ORDER BY ord")][:2] == ["en", "bn"]
        assert db.execute(
            "SELECT count(*) FROM words WHERE english = '' OR bangla IS NULL OR pron_bn IS NULL"
        ).fetchone()[0] == 0

    def test_1080_a_language_under_100_percent_is_held_back(self, russian, tmp_path, capsys):
        def one_blank(book):
            sheet = book["All Words"]
            sheet.cell(row=3, column=column_of(sheet, "Pronunciation (Russian)")).value = None

        partial = copy_with(russian / "German_B2_Tracker.xlsx", tmp_path, one_blank)
        tips = tmp_path / "tips.csv"
        tips.write_text("match_type,match,tip_ru,tags\npattern,^Tisch,Совет.,\n", encoding="utf-8")
        path = manifest(russian, {"German_B2_Tracker.xlsx": partial}, tips)

        db = build(tmp_path, path)
        err = capsys.readouterr().err
        assert (
            "language Russian (ru): meanings 180/180, pronunciations 179/180, examples 360/360"
            in err
        )
        assert "held back (German_B2_Tracker-one_blank.xlsx: pronunciations 29/30). --allow-partial ru builds it" in err
        assert "ru" not in [row[0] for row in db.execute("SELECT code FROM course_languages")]
        for table in ("word_meanings", "word_example_translations", "grammar_translations", "word_tips"):
            assert "ru" not in langs(db, table), table
        db.close()

        # For testing, once: built as it is.
        db = build(tmp_path, path, "--allow-partial", "ru")
        assert "partial, built for testing" in capsys.readouterr().err
        assert langs(db, "word_meanings")["ru"] == 180
        assert db.execute(
            "SELECT count(*) FROM word_meanings WHERE lang = 'ru' AND pronunciation IS NULL"
        ).fetchone()[0] == 1

    def test_1080_a_language_without_pronunciations_is_held_back(self, russian):
        # Meaning and pronunciation are required, even with no column for one.
        from excel_to_sqlite import derive
        from pipeline_steps import gate_languages

        sources = [read_workbook(russian / name) for name in BOOK_LEVELS]
        derive(sources)
        words = [word for source in sources for word in source.words]
        for word in words:
            del word.texts[("pronunciation", "ru")]
        shipped, report = gate_languages(words, [], {"en": {"meaning"}, "ru": {"meaning"}})
        assert [language.code for language in shipped] == ["en"]
        assert "language Russian (ru): meanings 180/180, pronunciations 0/180, held back" in report[1]

    @staticmethod
    def english_guides(russian, filled: int | None):
        """The fixture's words, the first [filled] (all, for None) with an English guide."""
        from excel_to_sqlite import derive

        sources = [read_workbook(russian / name) for name in BOOK_LEVELS]
        derive(sources)
        words = [word for source in sources for word in source.words]
        for word in words[:filled]:
            word.texts[("pronunciation", "en")] = "GOOT"
        return words

    @pytest.mark.parametrize("allow", [(), ("en",)])
    def test_1099_a_partial_english_guide_is_held_back_but_english_ships(self, russian, allow):
        from pipeline_steps import gate_languages

        words = self.english_guides(russian, 10)
        shipped, report = gate_languages(words, [], {"en": {"meaning", "pronunciation"}}, allow)
        assert [language.code for language in shipped] == ["en"]
        kept = sum(("pronunciation", "en") in word.texts for word in words)
        if allow:
            assert kept == 10 and "pronunciations 10/180, ships" in report[0]
        else:
            assert kept == 0
            assert "pronunciations 10/180, ships; pronunciations held back" in report[0]

    def test_1099_a_complete_english_guide_ships(self, russian):
        from pipeline_steps import gate_languages

        words = self.english_guides(russian, None)
        _, report = gate_languages(words, [], {"en": {"meaning", "pronunciation"}})
        assert all(("pronunciation", "en") in word.texts for word in words)
        assert report[0].endswith("pronunciations 180/180, ships")

    def test_1080_allow_partial_names_a_language_the_workbooks_carry(self, russian, tmp_path, capsys):
        out = tmp_path / "content.db"
        argv = ["--manifest", str(manifest(russian)), "--out", str(out), "--previous", str(tmp_path / "none")]
        assert build_main([*argv, "--allow-partial", "pl"]) == 1
        assert "--allow-partial pl: the workbooks carry no such language" in capsys.readouterr().err
