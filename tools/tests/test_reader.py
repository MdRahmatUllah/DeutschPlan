"""The workbook reader, against workbooks openpyxl actually wrote.

`tools/fixtures/make_workbooks.py` writes stand-ins with the shape
`content-pipeline.md` describes, because the four real trackers are not in the
repository. Testing the reader against mocks would prove nothing about headers
read by name, about a title row above the header, or about openpyxl.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from excel_to_sqlite import (  # noqa: E402
    GRAMMAR_HEADER_MAP,
    HEADER_MAP,
    REQUIRED_GRAMMAR_FIELDS,
    REQUIRED_WORD_FIELDS,
    PipelineError,
    Word,
    read_manifest,
    read_sources,
    read_workbook,
)
from fixtures.make_workbooks import (  # noqa: E402
    WORD_HEADERS as FIXTURE_HEADERS,
    BOOK_LEVELS,
    write_all,
)


# Column positions in the fixture, so a test that plants a cell says which
# column it means rather than counting on its fingers.
HEADER_INDEX = {name: i for i, name in enumerate(FIXTURE_HEADERS)}


@pytest.fixture(scope="module")
def books(tmp_path_factory) -> Path:
    directory = tmp_path_factory.mktemp("workbooks")
    write_all(directory)
    return directory


@pytest.fixture(scope="module")
def first_book(books: Path):
    return read_workbook(books / "German_B1_Tracker.xlsx")


def test_reads_every_sheet(first_book):
    assert first_book.words, "All Words produced no rows"
    assert first_book.grammar, "Grammar produced no rows"
    assert [c.name for c in first_book.categories] == ["Alltag", "Reisen", "Arbeit"]


def test_the_header_row_is_found_below_a_title_row(first_book):
    # The fixture puts a title in row 1, as the real workbooks do. A reader
    # that assumed row 1 was the header would find no columns at all.
    assert first_book.words[0].german
    assert first_book.words[0].english


def test_pipe_01_level_comes_from_the_cell_not_the_sheet(first_book):
    # The fixture writes "B1 (Phase 2)", which still means B1. The phase label
    # must not become the level, and must not stop the level being read.
    assert {w.level for w in first_book.words} == {"A1", "A2", "B1"}


DOC = (
    Path(__file__).resolve().parent.parent.parent
    / "docs"
    / "02-data"
    / "content-pipeline.md"
)


def _documented(heading: str) -> dict[str, tuple[str, bool]]:
    """Reads one header table out of content-pipeline.md.

    Parsed rather than restated: a table copied into a test proves only that
    someone copied it twice, and this one is the contract between the authors'
    column names and HEADER_MAP.
    """
    body = DOC.read_text(encoding="utf-8").split(heading, 1)[1]
    section = body.split(chr(10) + chr(10))[0]
    rows: dict[str, tuple[str, bool]] = {}
    for line in section.splitlines():
        cells = [c.strip() for c in line.split("|")]
        if len(cells) < 5 or cells[1].startswith("---") or cells[1] == "Header in":
            continue
        header, field, required = cells[1], cells[2].strip("`"), cells[3].lower()
        if not field or required not in {"yes", "no"}:
            continue
        rows[field] = (header, required == "yes")
    return rows


def test_the_doc_tables_were_found():
    # Everything below compares against these; an empty parse would make each
    # of them pass by having nothing to check.
    assert len(_documented("| Header in All Words |")) == 15
    assert len(_documented("| Header in Grammar |")) == 7


def test_HEADER_MAP_matches_the_documented_columns():
    documented = _documented("| Header in All Words |")
    assert set(HEADER_MAP) == set(documented)
    for field, (header, _) in documented.items():
        assert header in HEADER_MAP[field], (
            f"{field}: the doc calls the column {header!r}, HEADER_MAP has "
            f"{HEADER_MAP[field]!r}"
        )


def test_the_required_columns_are_the_ones_the_doc_marks_required():
    documented = _documented("| Header in All Words |")
    assert set(REQUIRED_WORD_FIELDS) == {
        field for field, (_, required) in documented.items() if required
    }


def test_the_grammar_map_matches_its_documented_columns():
    documented = _documented("| Header in Grammar |")
    assert set(GRAMMAR_HEADER_MAP) == set(documented)
    for field, (header, _) in documented.items():
        assert header in GRAMMAR_HEADER_MAP[field]
    assert set(REQUIRED_GRAMMAR_FIELDS) == {
        field for field, (_, required) in documented.items() if required
    }


def test_every_optional_column_is_mapped(first_book):
    # A header the map does not know is silently dropped. Checking that each
    # field was actually populated catches a spelling in HEADER_MAP that no
    # workbook carries, which the doc comparison above cannot see.
    fields = {
        f for f in HEADER_MAP if any(getattr(w, f, None) for w in first_book.words)
    }
    assert fields == set(HEADER_MAP), f"never populated: {set(HEADER_MAP) - fields}"


def test_freq_and_week_come_back_as_numbers(first_book):
    assert all(isinstance(w.freq, int) for w in first_book.words)
    assert all(isinstance(w.week, int) for w in first_book.words)


def test_a_formula_cell_survives_as_text(first_book):
    # PIPE-05. Storing it is this issue's job; warning about it is #46's.
    assert any(
        (w.collocations or "").startswith("=") for w in first_book.words
    ), "the fixture should carry a formula-looking cell"


def test_manifest_order_is_kept(books: Path, tmp_path: Path):
    manifest = tmp_path / "manifest.yaml"
    manifest.write_text(
        yaml.safe_dump(
            {
                "workbooks": [{"file": str(books / name)} for name in BOOK_LEVELS],
                "tips": "content/interference_tips.csv",
            }
        ),
        encoding="utf-8",
    )

    # The order is the fallback level order for grammar without a level, so it
    # is load-bearing rather than cosmetic.
    assert [p.name for p in read_manifest(manifest).workbooks] == list(BOOK_LEVELS)


def test_all_four_books_read(books: Path, tmp_path: Path):
    manifest = tmp_path / "manifest.yaml"
    manifest.write_text(
        yaml.safe_dump(
            {"workbooks": [{"file": str(books / name)} for name in BOOK_LEVELS]}
        ),
        encoding="utf-8",
    )

    sources = read_sources(read_manifest(manifest))
    assert {w.level for source in sources for w in source.words} == {
        "A1",
        "A2",
        "B1",
        "B2",
        "C1",
        "C2",
    }


def test_the_repository_manifest_parses():
    # The real one, so a typo in content/manifest.yaml fails here rather than
    # on someone's machine at `make content`.
    repo_root = Path(__file__).resolve().parent.parent.parent
    manifest = read_manifest(repo_root / "content" / "manifest.yaml")
    assert [p.name for p in manifest.workbooks] == list(BOOK_LEVELS)
    assert manifest.tips is not None and manifest.tips.name == "interference_tips.csv"


class TestFailures:
    """Every one of these must say what to do, not just that something broke."""

    def test_a_missing_workbook_names_the_file_and_the_manifest(self, tmp_path: Path):
        manifest = tmp_path / "manifest.yaml"
        manifest.write_text(
            yaml.safe_dump({"workbooks": [{"file": str(tmp_path / "Nope.xlsx")}]}),
            encoding="utf-8",
        )
        with pytest.raises(PipelineError, match=r"Nope\.xlsx is missing.*manifest"):
            read_sources(read_manifest(manifest))

    def test_a_missing_sheet_lists_the_sheets_that_are_there(
        self, books: Path, tmp_path: Path
    ):
        path = _without_sheet(books / "German_B2_Tracker.xlsx", "Grammar", tmp_path)
        with pytest.raises(PipelineError, match=r"no 'Grammar' sheet.*All Words"):
            read_workbook(path)

    def test_a_renamed_required_header_points_at_HEADER_MAP(
        self, books: Path, tmp_path: Path
    ):
        path = _rename_header(
            books / "German_C1_Tracker.xlsx", "German", "Deutsch", tmp_path
        )
        with pytest.raises(PipelineError, match=r"HEADER_MAP"):
            read_workbook(path)

    def test_a_renamed_optional_header_is_not_fatal(
        self, books: Path, tmp_path: Path
    ):
        # Optional means optional: losing the Bangla column must not stop a
        # build the evening before a release.
        path = _rename_header(
            books / "German_C2_Tracker.xlsx", "Bangla meaning", "Bangla?", tmp_path
        )
        source = read_workbook(path)
        assert source.words
        assert all(w.bangla is None for w in source.words)

    def test_a_row_missing_a_required_cell_names_the_row(
        self, books: Path, tmp_path: Path
    ):
        path = _blank_cell(books / "German_C2_Tracker.xlsx", "English", tmp_path)
        with pytest.raises(PipelineError, match=r"row \d+: missing English"):
            read_workbook(path)

    def test_a_workbook_with_no_category_tab_says_so(
        self, books: Path, tmp_path: Path
    ):
        def drop_categories(book):
            for name in list(book.sheetnames):
                if name.startswith("C-"):
                    book.remove(book[name])

        path = _copy_with(books / "German_B2_Tracker.xlsx", tmp_path, drop_categories)
        with pytest.raises(PipelineError, match=r"no 'C-…' category tab"):
            read_workbook(path)

    def test_a_divider_row_is_skipped_rather_than_fatal(
        self, books: Path, tmp_path: Path
    ):
        # A banner row carrying only the level, which is how these sheets
        # separate one block from the next. It has a Level and no word, so a
        # rule that only skips wholly blank rows would stop the build on it and
        # send the author to delete a row that looks deliberate to them.
        def add_divider(book):
            sheet = book["All Words"]
            row = [None] * len(HEADER_MAP)
            row[HEADER_INDEX["Level"]] = "C1"
            sheet.append(row)

        path = _copy_with(books / "German_C1_Tracker.xlsx", tmp_path, add_divider)
        assert read_workbook(path).words

    def test_a_word_row_missing_its_level_is_still_fatal(
        self, books: Path, tmp_path: Path
    ):
        # The other half of the rule above: a row that names a word and then
        # omits its level would be filed under the wrong step silently.
        path = _blank_cell(books / "German_C1_Tracker.xlsx", "Level", tmp_path)
        with pytest.raises(PipelineError, match=r"row \d+: missing Level"):
            read_workbook(path)

    def test_a_manifest_with_no_workbooks_says_so(self, tmp_path: Path):
        manifest = tmp_path / "manifest.yaml"
        manifest.write_text("workbooks: []\n", encoding="utf-8")
        with pytest.raises(PipelineError, match="lists no workbooks"):
            read_manifest(manifest)


def test_blank_padding_rows_are_skipped(books: Path, tmp_path: Path):
    path = tmp_path / "padded.xlsx"
    book = load_workbook(books / "German_C1_Tracker.xlsx")
    sheet = book["All Words"]
    before = sheet.max_row
    for _ in range(20):
        sheet.append([None] * len(HEADER_MAP))
    book.save(path)

    source = read_workbook(path)
    # Two title/header rows above the data.
    assert len(source.words) == before - 2


def _copy_with(path: Path, tmp_path: Path, mutate) -> Path:
    book = load_workbook(path)
    mutate(book)
    out = tmp_path / f"{path.stem}-{mutate.__name__}.xlsx"
    book.save(out)
    return out


def _without_sheet(path: Path, name: str, tmp_path: Path) -> Path:
    def drop(book):
        book.remove(book[name])

    return _copy_with(path, tmp_path, drop)


def _rename_header(path: Path, old: str, new: str, tmp_path: Path) -> Path:
    def rename(book):
        for row in book["All Words"].iter_rows(min_row=1, max_row=5):
            for cell in row:
                if cell.value == old:
                    cell.value = new

    return _copy_with(path, tmp_path, rename)


def _blank_cell(path: Path, header: str, tmp_path: Path) -> Path:
    def blank(book):
        sheet = book["All Words"]
        column = None
        for row in sheet.iter_rows(min_row=1, max_row=5):
            for cell in row:
                if cell.value == header:
                    column = cell.column
            if column:
                break
        # The first data row after the title and header rows.
        sheet.cell(row=3, column=column).value = None

    return _copy_with(path, tmp_path, blank)


def test_word_is_a_plain_record():
    # Nothing downstream should need the reader's own types to construct one.
    assert Word(source_file="x", row=2, german="Haus", english="house", level="A1")
