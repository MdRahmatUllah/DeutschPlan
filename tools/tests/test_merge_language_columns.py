"""#1082-#1084: tools/merge_language_columns.py writes a meaning language's
columns into a copy of a workbook, and nothing else."""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font
from openpyxl.worksheet.formula import ArrayFormula

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from merge_language_columns import MergeError, chart_parts, check, merge, put_parts  # noqa: E402

MARK = b"<!-- the original chart -->"
WORD_HEADERS = ["ID", "Article", "German", "POS", "Pronunciation (Bangla)", "English",
                "Examples (DE)", "Examples (EN)"]


def fixture(path: Path) -> Path:
    """A tracker in miniature: a Dashboard chart, All Words and Grammar with
    header row 4, a grouped column width past the last column (as B2-C2) and
    an array formula (as Add Words)."""
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    for i in range(1, 5):
        dash.append([i, i * 10])
    dash["F1"] = ArrayFormula("F1", "=SUM(B1:B4*2)")  # compares by identity, not by its text
    chart = LineChart()
    chart.add_data(Reference(dash, min_col=2, min_row=1, max_row=4))
    dash.add_chart(chart, "D2")
    words = wb.create_sheet("All Words")
    words["A1"] = "All Words — master list"
    words.append([])
    words.append([])
    words.append(WORD_HEADERS)
    for c in words[4]:
        c.font = Font(bold=True, color="FFFFFF")
    words.append([1, None, "ich", "pron", "ইশ", "I", "Ich komme.\nIch gehe.", "I come.\nI go."])
    words.append([2, "das", "Buch", "noun", "বুখ", "book", "Das Buch ist neu.", "The book is new."])
    words.append([3, None, "Bank", "noun", "বাংক", "bank", "Die Bank ist zu.", "The bank is closed."])
    words.column_dimensions.group("I", "P", hidden=False)
    words.column_dimensions["I"].width = 8.71
    grammar = wb.create_sheet("Grammar")
    grammar.append([])
    grammar.append([])
    grammar.append([])
    grammar.append(["#", "Week", "Level", "Topic", "Rule", "Example (DE)", "Example (EN)", "Watch out"])
    grammar.append([1, 1, "A1", "sein", "ich bin", "Ich bin hier.", "I am here.", "sind"])
    grammar.append([2, 2, "A1", "haben", "ich habe", "Ich habe Zeit.", "I have time.", "hat"])
    wb.save(path)
    # openpyxl never writes an XML comment: the chart it rewrites on save loses this
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["xl/charts/chart1.xml"] += MARK
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return path


RU = {
    "language": "Russian",
    "words": [
        {"key": ["ich", "pron", "I"], "Meaning": "я", "Pronunciation": "ихь",
         "Examples": ["Я иду.", "Я ухожу."]},
        {"key": ["Buch", "noun", "book"], "Meaning": "книга", "Pronunciation": "бух",
         "Examples": ["Книга новая."]},
    ],
    "grammar": [{"key": ["A1", "sein"], "Topic": "sein", "Rule": "ich bin",
                 "Example": "Я здесь.", "Watch out": "sind"}],
}
#: the columns RU writes
WRITTEN = {"All Words": {9, 10, 11}, "Grammar": {9, 10, 11, 12}}
EN = {"language": "English", "words": [{"key": ["ich", "pron", "I"], "Pronunciation": "IHY"}]}


@pytest.fixture
def src(tmp_path: Path) -> Path:
    return fixture(tmp_path / "German_T1_Tracker.xlsx")


def headers(ws, row: int = 4) -> list:
    return [c.value for c in ws[row]]


def test_the_columns_go_after_the_last_one_and_the_input_is_untouched(src, tmp_path):
    before = src.read_bytes()
    out = tmp_path / "out" / src.name
    merge(src, out, [EN, RU])
    assert src.read_bytes() == before
    wb = load_workbook(out)
    words = wb["All Words"]
    assert headers(words) == WORD_HEADERS + [
        "Pronunciation (English)", "Meaning (Russian)", "Pronunciation (Russian)", "Examples (Russian)"]
    assert [c.value for c in words[5]][8:] == ["IHY", "я", "ихь", "Я иду.\nЯ ухожу."]
    # a word the staging lacks stays blank
    assert [c.value for c in words[7]][8:] == [None, None, None, None]
    assert words.cell(4, 10).font.bold and words.cell(4, 10).font.color.rgb == "00FFFFFF"
    assert words.column_dimensions["L"].width == 38
    # the split group keeps its width on the columns the merge did not touch
    assert words.column_dimensions["M"].width == 8.71
    assert headers(wb["Grammar"])[8:] == [
        "Topic (Russian)", "Rule (Russian)", "Example (Russian)", "Watch out (Russian)"]
    assert [c.value for c in wb["Grammar"][6]][8:] == [None] * 4


def test_the_chart_is_the_original(src, tmp_path):
    out = tmp_path / "out.xlsx"
    merge(src, out, [RU])
    with zipfile.ZipFile(out) as z:
        assert z.read("xl/charts/chart1.xml").endswith(MARK)


def test_a_chart_openpyxl_drops_is_put_back(src, tmp_path, monkeypatch):
    monkeypatch.setattr("openpyxl.reader.excel.find_images", lambda *a: ([], []))
    out = tmp_path / "out.xlsx"
    merge(src, out, [RU])
    with zipfile.ZipFile(out) as z:
        assert z.read("xl/charts/chart1.xml").endswith(MARK)


def test_merging_again_rewrites_the_columns_in_place(src, tmp_path):
    first, second = tmp_path / "1.xlsx", tmp_path / "2.xlsx"
    merge(src, first, [RU])
    fixed = {**RU, "words": [{**RU["words"][0], "Meaning": "я (исправлено)"}]}
    merge(first, second, [fixed])
    words = load_workbook(second)["All Words"]
    assert headers(words).count("Meaning (Russian)") == 1
    assert words.cell(5, 9).value == "я (исправлено)"
    assert words.cell(6, 9).value is None  # the staging is the whole column


def test_the_check_catches_a_changed_cell(src, tmp_path):
    out = tmp_path / "out.xlsx"
    merge(src, out, [RU])
    wb = load_workbook(out)
    wb["All Words"]["F6"] = "a book"
    wb.save(out)
    with pytest.raises(MergeError, match="All Words!F6"):
        check(src, out, WRITTEN)


def test_the_check_catches_a_changed_chart(src, tmp_path):
    out = tmp_path / "out.xlsx"
    merge(src, out, [RU])
    put_parts(out, {n: d.replace(MARK, b"") for n, d in chart_parts(out).items()})
    with pytest.raises(MergeError, match="chart parts"):
        check(src, out, WRITTEN)


@pytest.mark.parametrize("staging, why", [
    ({**RU, "words": [{**RU["words"][0], "key": ["ich", "pron", "me"]}]}, "no row has the key"),
    ({**RU, "words": [{**RU["words"][0], "Examples": ["Я иду."]}]}, "1 example lines"),
    ({**RU, "words": [RU["words"][0], RU["words"][0]]}, "staged twice"),
    ({**RU, "words": [RU["words"][0], {"key": ["Buch", "noun", "book"], "Meaning": "книга"}]},
     "fields"),
    ({**RU, "grammar": [{"key": ["A1", "sein"], "Topik": "sein"}]}, "unknown fields"),
])
def test_a_bad_staging_stops_the_merge_and_writes_nothing(src, tmp_path, staging, why):
    out = tmp_path / "out.xlsx"
    with pytest.raises(MergeError, match=why):
        merge(src, out, [staging])
    assert not out.exists()


def test_it_never_writes_its_input(src):
    with pytest.raises(MergeError, match="never writes its input"):
        merge(src, src, [RU])


def test_a_key_on_two_rows_stops_the_merge(src, tmp_path):
    wb = load_workbook(src)
    wb["All Words"].append([4, None, "ich", "pron", "ইশ", "I", "Ich.", "I."])
    wb.save(src)
    with pytest.raises(MergeError, match="the same key"):
        merge(src, tmp_path / "out.xlsx", [RU])
