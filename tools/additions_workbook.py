"""#1257: writes `data/German_Everyday_Additions.xlsx`, the course's additions
workbook, from the reviewed `content/additions/*.yaml`.

The six trackers are the owner's study files, each row wired into its week,
category and Dashboard sheets, and no tool writes them (#545). Words the
course adds go here instead (the owner, 2026-10-03): this workbook is built
from the YAML, so it is regenerated, never edited. Each word carries its own
Level and Week, so it lands in its step, and takes its week's place there
(`derive`'s sort), as a tracker row would.

    python tools/additions_workbook.py           # data/German_Everyday_Additions.xlsx
    python tools/additions_workbook.py --check   # the YAML only
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook

from pipeline_steps import LEVELS

ROOT = Path(__file__).resolve().parents[1]
SOURCES = ROOT / "content" / "additions"
OUT = ROOT / "data" / "German_Everyday_Additions.xlsx"
CATEGORY_NAMES = ROOT / "content" / "category_names.csv"

#: A meaning language's code in the YAML, and its name in the headers.
LANGUAGES = {"ru": "Russian", "pl": "Polish"}

#: The YAML's fields, in the trackers' header names (`HEADER_MAP`).
FIELDS = {
    "article": "Article",
    "german": "German",
    "forms": "Plural / forms",
    "pos": "POS",
    "pron_bn": "Pronunciation (Bangla)",
    "english": "English",
    "bangla": "Bangla meaning",
    "freq": "Freq",
    "level": "Level",
    "category": "Category",
    "week": "Week",
    "examples_de": "Examples (DE)",
    "examples_en": "Examples (EN)",
    "pron_en": "Pronunciation (English)",
}
REQUIRED = tuple(f for f in FIELDS if f not in ("article", "forms"))
LANGUAGE_PARTS = {"meaning": "Meaning", "pron": "Pronunciation", "examples": "Examples"}

WORD_HEADERS = [
    *FIELDS.values(),
    *[f"{part} ({name})" for name in LANGUAGES.values() for part in LANGUAGE_PARTS.values()],
]
GRAMMAR_HEADERS = [
    "Week", "Level", "Topic", "Rule", "Example (DE)", "Example (EN)", "Watch out",
    *[
        f"{part} ({name})"
        for name in LANGUAGES.values()
        for part in ("Topic", "Rule", "Example", "Watch out")
    ],
]


def read_entries(sources: Path = SOURCES) -> list[dict]:
    entries: list[dict] = []
    for path in sorted(sources.glob("*.yaml")):
        for entry in yaml.safe_load(path.read_text(encoding="utf-8"))["words"]:
            entries.append({**entry, "_file": path.name})
    return entries


def category_names(path: Path = CATEGORY_NAMES) -> set[str]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return {row["name"] for row in csv.DictReader(handle)}


def problems(entries: list[dict], categories: set[str]) -> list[str]:
    """What stops the workbook: a missing field, an unknown level or category,
    example lines that don't pair, a word twice."""
    found: list[str] = []
    seen: set[tuple] = set()
    for e in entries:
        name = f"{e['_file']}: {e.get('german')!r}"
        found += [f"{name}: no {f}" for f in REQUIRED if e.get(f) in (None, "")]
        if e.get("level") not in LEVELS:
            found.append(f"{name}: level {e.get('level')!r}")
        if e.get("category") not in categories:
            found.append(f"{name}: category {e.get('category')!r} has no names in {CATEGORY_NAMES.name}")
        if not isinstance(e.get("week"), int) or e["week"] < 1:
            found.append(f"{name}: week {e.get('week')!r}")
        if not isinstance(e.get("freq"), int) or not 1 <= e["freq"] <= 5:
            found.append(f"{name}: freq {e.get('freq')!r}")
        lines = len(str(e.get("examples_de", "")).split("\n"))
        if len(str(e.get("examples_en", "")).split("\n")) != lines:
            found.append(f"{name}: the English examples don't pair with the German")
        for code in LANGUAGES:
            language = e.get(code) or {}
            found += [f"{name}: no {code} {part}" for part in LANGUAGE_PARTS if not language.get(part)]
            if language.get("examples") and len(language["examples"].split("\n")) != lines:
                found.append(f"{name}: the {code} examples don't pair with the German")
        key = (e.get("level"), e.get("german"), e.get("pos"), e.get("english"))
        if key in seen:
            found.append(f"{name}: twice")
        seen.add(key)
    return found


def row(entry: dict) -> list:
    return [
        *[entry.get(field) for field in FIELDS],
        *[entry[code][part] for code in LANGUAGES for part in LANGUAGE_PARTS],
    ]


def write(entries: list[dict], out: Path = OUT) -> None:
    book = Workbook()
    words = book.active
    words.title = "All Words"
    words.append(WORD_HEADERS)
    for entry in entries:
        words.append(row(entry))
    book.create_sheet("Grammar").append(GRAMMAR_HEADERS)
    # The build needs a category tab in every workbook: the ones its words use.
    for category in sorted({e["category"] for e in entries}):
        book.create_sheet(f"C-{category}")
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--check", action="store_true", help="check the YAML, write nothing")
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args(argv)
    entries = read_entries()
    found = problems(entries, category_names())
    for problem in found:
        print(problem, file=sys.stderr)
    if found:
        return 1
    if not args.check:
        write(entries, args.out)
        print(f"{len(entries)} words -> {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
